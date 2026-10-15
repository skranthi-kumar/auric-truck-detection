"""report/REPORT.md -> report/report.docx (headings, paragraphs, nested lists, tables, figures).

Handles the Markdown subset used in REPORT.md: #/## headings, **bold**, `code`, -/1. lists (2-space nesting),
pipe tables, ![alt](path) and <img src=... style="width:NN%"> figures.
Usage: python src/build_report_docx.py
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

REP = Path(__file__).resolve().parents[1] / "report"
TEXT_W = 17.0  # cm, A4 minus margins
INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")


def runs(p, text):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            p.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = p.add_run(part[1:-1])
            r.font.name = "Menlo"
            r.font.size = Pt(8.5)
        else:
            p.add_run(part)


def shade(cell, hexcolor):
    tc = cell._tc.get_or_add_tcPr()
    s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear"), s.set(qn("w:color"), "auto"), s.set(qn("w:fill"), hexcolor)
    tc.append(s)


def table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [cells[0]] + cells[2:]  # drop the |---| separator row
    t = doc.add_table(rows=len(cells), cols=len(cells[0]))
    t.style = "Table Grid"
    for i, r in enumerate(cells):
        for j, v in enumerate(r):
            c = t.cell(i, j)
            c.paragraphs[0].text = ""
            runs(c.paragraphs[0], v)
            for run in c.paragraphs[0].runs:
                run.font.size = Pt(8.5)
                if i == 0:
                    run.bold = True
            if i == 0:
                shade(c, "F1F1EE")
    doc.add_paragraph()


def figure(doc, path, width_pct=100):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(REP / path), width=Cm(TEXT_W * width_pct / 100))


def main():
    import argparse
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", default="REPORT.md")
    ap.add_argument("--out", default="report.docx")
    a = ap.parse_args()
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(10)
    st.paragraph_format.space_after = Pt(3)
    for name, size in (("Heading 1", 17), ("Heading 2", 13)):
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].font.color.rgb = RGBColor(0x11, 0x11, 0x11)

    lines = (REP / a.src).read_text().splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        ind = len(line) - len(line.lstrip())
        if not s:
            i += 1
            continue
        if s.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            table(doc, block)
            continue
        m = re.match(r"!\[[^\]]*\]\(([^)]+)\)", s) or re.match(r'<img src="([^"]+)"(?: style="width:(\d+)%")?>', s)
        if m:
            figure(doc, m.group(1), int(m.group(2)) if m.lastindex and m.lastindex > 1 and m.group(2) else 100)
        elif s.startswith("# "):
            doc.add_heading(s[2:], level=1)
        elif s.startswith("## "):
            doc.add_heading(s[3:], level=2)
        elif s.startswith("### "):
            doc.add_heading(s[4:], level=3)
        elif s.startswith("- ") or re.match(r"\d+\. ", s):
            numbered = not s.startswith("- ")
            level = min(ind // 2, 2)
            style = ("List Number" if numbered else "List Bullet") + ("" if level == 0 else f" {level + 1}")
            p = doc.add_paragraph(style=style)
            p.paragraph_format.space_after = Pt(1)
            runs(p, s[2:] if not numbered else s.split(". ", 1)[1])
        elif s.startswith("*") and s.endswith("*") and not s.startswith("**"):
            p = doc.add_paragraph()
            r = p.add_run(s.strip("*"))
            r.italic = True
        else:
            runs(doc.add_paragraph(), s)
        i += 1
    out = REP / a.out
    doc.save(out)
    print("wrote", out)


if __name__ == "__main__":
    main()
